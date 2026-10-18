"""
Generate comprehensive, professional academic project report in DOCX format
for TalentLens AI — AI Resume Screening System.
"""

from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

BASE_DIR = Path(__file__).resolve().parent.parent
SUBMISSION_DIR = BASE_DIR / "submission_docs"
IMAGES_DIR = SUBMISSION_DIR / "images"
DOCX_OUT = SUBMISSION_DIR / "TalentLens_AI_Project_Report.docx"


def set_cell_background(cell, hex_color):
    """Sets background shading of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets inner padding for a table cell in twips."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def add_callout_box(doc, text, title="ETHICAL AI NOTICE"):
    """Adds a styled callout box with a colored border."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0FDF4")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    # Set left border thick green
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="10B981"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_title = p.add_run(f"⚖️ {title}: ")
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(6, 95, 70)
    run_title.font.size = Pt(10)

    run_text = p.add_run(text)
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = RGBColor(55, 65, 81)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def build_docx_report():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Default Style Configurations
    style_normal = doc.styles["Normal"]
    style_normal.font.name = "Calibri"
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(31, 41, 55)
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(6)

    # ---------------------------------------------------------
    # 1. TITLE / COVER PAGE
    # ---------------------------------------------------------
    p_title_space = doc.add_paragraph()
    p_title_space.paragraph_format.space_before = Pt(36)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("TalentLens AI — AI Resume Screening System")
    run_title.bold = True
    run_title.font.size = Pt(26)
    run_title.font.color.rgb = RGBColor(30, 58, 138)  # Deep Navy

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("An Explainable, Local, and Privacy-Preserving Candidate Screening Pipeline")
    run_sub.font.size = Pt(14)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(75, 85, 99)

    doc.add_paragraph().paragraph_format.space_before = Pt(40)

    # Submission Meta Table
    meta_tbl = doc.add_table(rows=6, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_info = [
        ("Course / Degree:", "[ENTER_DEGREE_PROGRAM, e.g., B.Tech in Computer Science & Engineering]"),
        ("Candidate / Author:", "[STUDENT_NAME] (Roll No: [ROLL_NUMBER])"),
        ("Supervisor / Guide:", "[SUPERVISOR_NAME / FACULTY_GUIDE]"),
        ("Department:", "[DEPARTMENT_NAME, e.g., Department of Computer Science & Engineering]"),
        ("Institution / University:", "[COLLEGE_OR_UNIVERSITY_NAME]"),
        ("Source Code Repository:", "https://github.com/Rajendra13188/TalentLens-AI"),
    ]
    for idx, (label, val) in enumerate(meta_info):
        r = meta_tbl.rows[idx]
        c1, c2 = r.cells[0], r.cells[1]
        c1.width = Inches(2.2)
        c2.width = Inches(4.3)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(label)
        r1.bold = True
        r1.font.color.rgb = RGBColor(30, 58, 138)
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(val)
        if "[" in val:
            r2.italic = True
            r2.font.color.rgb = RGBColor(180, 83, 9)  # Highlight placeholder in amber
        else:
            r2.font.color.rgb = RGBColor(31, 41, 55)

    doc.add_paragraph().paragraph_format.space_before = Pt(80)

    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_date = p_date.add_run("Academic Project Report • 2026")
    r_date.font.size = Pt(10)
    r_date.font.color.rgb = RGBColor(107, 114, 128)

    doc.add_page_break()

    # Helper function for headings
    def add_sec_heading(title, level=1):
        h = doc.add_heading(title, level=level)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        if level == 1:
            h.runs[0].font.color.rgb = RGBColor(30, 58, 138)
            h.runs[0].font.size = Pt(18)
        elif level == 2:
            h.runs[0].font.color.rgb = RGBColor(37, 99, 235)
            h.runs[0].font.size = Pt(14)
        return h

    # ---------------------------------------------------------
    # 2. ABSTRACT
    # ---------------------------------------------------------
    add_sec_heading("1. Abstract")
    doc.add_paragraph(
        "Modern technical hiring is inundated with hundreds of applicants per role, compelling organizations to adopt "
        "automated Applicant Tracking Systems (ATS). However, commercial ATS solutions often operate as opaque 'black boxes' "
        "or rely on expensive, proprietary cloud APIs that expose confidential candidate resumes to third parties. Naive keyword "
        "matchers further exacerbate the issue by producing arbitrary binary rejections without explainability or context."
    )
    doc.add_paragraph(
        "This project presents TalentLens AI, an explainable, 100% local artificial intelligence resume screening system "
        "engineered in Python using Streamlit, scikit-learn, PyMuPDF, and pandas. TalentLens AI ingests candidate PDF resumes "
        "directly into memory, extracts clean digital text, gracefully isolates unreadable or scanned documents, and evaluates "
        "candidates against job descriptions using a dual-engine architecture: (1) an n-gram TF-IDF cosine similarity model "
        "capturing contextual vocabulary overlap, and (2) a regex-based skills taxonomy engine spanning seven technical domains. "
        "An explainable scoring formula derives a deterministic score from 0 to 100 with an exact mathematical breakdown. "
        "The system enforces strict fair hiring principles, guarantees zero persistent resume storage, and provides interactive "
        "leaderboards, candidate deep-dive inspectors, and skill gap analytics."
    )

    # ---------------------------------------------------------
    # 3. PROBLEM STATEMENT & OBJECTIVES
    # ---------------------------------------------------------
    add_sec_heading("2. Problem Statement & Objectives")
    add_sec_heading("2.1 Problem Statement", level=2)
    doc.add_paragraph(
        "Recruitment teams face three critical challenges when screening technical resumes:\n"
        "1. Opacity and Lack of Explainability: Deep neural embeddings and closed-source screening tools generate arbitrary ranking scores "
        "without disclosing why a candidate was ranked favorably or discarded.\n"
        "2. Privacy and Security Vulnerabilities: Cloud-hosted resume parsers upload sensitive personal identifiable information (PII) "
        "to third-party servers, posing GDPR and privacy compliance risks.\n"
        "3. Brittle Document Ingestion: Many parsers crash when encountering corrupted, blank, or scanned image-only PDFs, halting "
        "batch screening pipelines."
    )

    add_sec_heading("2.2 Project Objectives", level=2)
    doc.add_paragraph(
        "The primary objectives of TalentLens AI are:\n"
        "• Local and Sovereign Execution: Implement a fully functional resume screening system that executes entirely offline on standard Windows hardware without requiring paid APIs or internet access.\n"
        "• Mathematically Explainable Scoring: Formulate an interpretable 0–100 matching equation decomposing scores into TF-IDF cosine similarity, direct job skill overlap, and domain breadth bonuses.\n"
        "• Fault-Tolerant PDF Extraction: Utilize PyMuPDF to extract digital text while automatically classifying corrupted, empty, and scanned image documents without pipeline interruption.\n"
        "• Ethical AI Governance: Eliminate protected demographic characteristics from ranking algorithms and display prominent disclaimers explaining that textual similarity is not a probability of human job performance.\n"
        "• Dynamic Recruiter Interface: Build an interactive Streamlit dashboard supporting custom job descriptions, presets, individual candidate inspectors, pool skill-gap histograms, and CSV data export."
    )

    # ---------------------------------------------------------
    # 4. EXISTING VS PROPOSED SYSTEM
    # ---------------------------------------------------------
    add_sec_heading("3. Existing Systems vs. Proposed TalentLens AI")
    doc.add_paragraph(
        "A rigorous comparative analysis reveals substantial advantages of TalentLens AI over legacy ATS approaches:"
    )

    comp_tbl = doc.add_table(rows=7, cols=3)
    comp_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Feature / Dimension", "Traditional / Cloud ATS", "TalentLens AI (Proposed)"]
    hdr_row = comp_tbl.rows[0]
    for idx, name in enumerate(headers):
        cell = hdr_row.cells[idx]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        r = p.add_run(name)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    comp_data = [
        ("Deployment & Cost", "Proprietary cloud SaaS; costly per-resume API fees", "100% Free, local open-source Python execution"),
        ("Data Privacy & Storage", "Uploaded resumes retained on external vendor servers", "In-memory processing only; zero persistent disk storage"),
        ("Scoring Explainability", "Opaque black-box AI scores or boolean keywords", "Exact mathematical breakdown formula (0–100) displayed in UI"),
        ("Scanned/Corrupt PDFs", "Throws unhandled exceptions or silently fails", "Robust PyMuPDF error catching; identifies scanned/empty files"),
        ("Skill Taxonomy", "Hard-coded or inflexible third-party dictionaries", "7-domain taxonomy with symbol lookarounds and UI custom skill additions"),
        ("Fairness & Ethics", "Often biased by demographics, names, or formatting", "Demographic neutrality; mandatory caveats that absent keywords ≠ lack of skill")
    ]
    for row_idx, (dim, trad, proposed) in enumerate(comp_data, start=1):
        row = comp_tbl.rows[row_idx]
        bg = "F9FAFB" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate([dim, trad, proposed]):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 80, 80)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            if col_idx == 0:
                r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ---------------------------------------------------------
    # 5. TECHNOLOGY STACK
    # ---------------------------------------------------------
    add_sec_heading("4. Technology Stack & Design Principles")
    doc.add_paragraph(
        "TalentLens AI is engineered using modern, robust, and lightweight Python libraries:\n"
        "• Python 3.13: Modern runtime environment providing strong typing, performance, and process isolation.\n"
        "• Streamlit (1.65.0): High-performance reactive web framework enabling rich UI dashboards, interactive sliders, data tables, and dynamic state management.\n"
        "• scikit-learn (1.9.1): Industrial-grade machine learning library supplying TfidfVectorizer and cosine_similarity for numerical vector comparisons.\n"
        "• PyMuPDF / pymupdf (1.28.2): High-speed C-backed MuPDF document parser for digital text extraction, stream validation, and image block counting.\n"
        "• pandas (3.0.6) & numpy (2.5.4): High-performance tabular data structures powering leaderboard filtering, metric aggregations, and CSV generation.\n"
        "• pytest (9.1.1): Comprehensive automated test framework for regression testing, edge case verification, and headless UI workflow execution."
    )

    # ---------------------------------------------------------
    # 6. SYSTEM ARCHITECTURE & WORKFLOW
    # ---------------------------------------------------------
    add_sec_heading("5. System Architecture & Workflow")
    doc.add_paragraph(
        "The architecture is organized into decoupled layers: Document Ingestion, Natural Language Processing, "
        "Scoring Engine, and Presentation Dashboard."
    )

    arch_img_path = IMAGES_DIR / "architecture_diagram.png"
    if arch_img_path.exists():
        doc.add_picture(str(arch_img_path), width=Inches(6.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figure 1: TalentLens AI System Architecture & Pipeline Workflow")
        r_cap.italic = True
        r_cap.font.size = Pt(9.5)
        r_cap.font.color.rgb = RGBColor(107, 114, 128)

    doc.add_paragraph(
        "Pipeline Execution Stages:\n"
        "1. Ingestion: The user inputs a job description (or selects a preset) and provides PDF resumes via file upload or synthetic demo loading.\n"
        "2. PyMuPDF Processing: In-memory byte streams are inspected. If character count is <40 and images exist, the file is classified as a scanned image. Blank or corrupt files are tagged with error states.\n"
        "3. TF-IDF Representation: Clean resume texts and the job description are transformed into an n-gram (1-2) TF-IDF matrix with sublinear scaling.\n"
        "4. Skill Taxonomy Comparison: Regex lookaround engines scan for canonical technical skills and aliases, computing matched, missing, and additional skill sets.\n"
        "5. Explainable Scoring: The scoring engine combines the metrics into a 0–100 score, orders candidates, and renders the interactive dashboard."
    )

    # ---------------------------------------------------------
    # 7. ALGORITHMS & MATHEMATICAL FORMULATION
    # ---------------------------------------------------------
    add_sec_heading("6. Algorithms & Mathematical Formulation")
    doc.add_paragraph(
        "Candidate alignment is calculated using a transparent mathematical model rather than uninterpretable neural weights."
    )

    add_sec_heading("6.1 Term Frequency - Inverse Document Frequency (TF-IDF)", level=2)
    doc.add_paragraph(
        "TF-IDF quantifies the lexical relevance of vocabulary terms between the job description and candidate documents. "
        "Sublinear term frequency scaling is applied to prevent repetitive term stuffing:"
    )
    doc.add_paragraph(
        "Term Frequency:  tf(t, d) = 1 + ln(count(t, d))  if count > 0, else 0\n"
        "Inverse Document Frequency:  idf(t, D) = ln((1 + |D|) / (1 + df(t))) + 1\n"
        "TF-IDF Weight:  w(t, d) = tf(t, d) × idf(t, D)"
    )
    doc.add_paragraph(
        "Cosine similarity measures the angular alignment between the normalized job description vector v_JD and resume vector v_Resume in n-dimensional term space:\n"
        "Cosine Similarity:  cos(v_JD, v_Resume) = (v_JD • v_Resume) / (||v_JD|| × ||v_Resume||)\n"
        "TF-IDF Score Component:  S_tfidf = np.clip(Cosine Similarity × 100.0, 0.0, 100.0)"
    )

    add_sec_heading("6.2 Direct Skill Coverage Score", level=2)
    doc.add_paragraph(
        "The skills taxonomy identifies required skills in the job description (Skills_JD) and candidate resume (Skills_Resume). "
        "Direct skill match ratio calculates the fraction of mandatory skills satisfied:\n"
        "S_skills = (|Skills_JD ∩ Skills_Resume| / |Skills_JD|) × 100.0  (if |Skills_JD| > 0, else S_tfidf)"
    )

    add_sec_heading("6.3 Domain Breadth Bonus", level=2)
    doc.add_paragraph(
        "To reward well-rounded candidates possessing complementary technical proficiencies outside the immediate job description, "
        "a breadth bonus is computed (capped at 10.0 points):\n"
        "B_breadth = min(10.0, |Skills_Resume \\ Skills_JD| × 1.5)"
    )

    add_sec_heading("6.4 Final Weighted Scoring Equation", level=2)
    doc.add_paragraph(
        "The overall matching score is synthesized using non-negative weights that sum to 1.00 (default: 40% skills, 50% TF-IDF, 10% breadth):\n"
        "Final Score = min(100.0, max(0.0, (w_skills × S_skills) + (w_tfidf × S_tfidf) + (w_breadth × B_breadth)))\n\n"
        "Example Candidate Formula Breakdown Displayed in UI:\n"
        "(17.5 × 0.50) + (100.0 × 0.40) + (1.5 × 0.10) = 48.9 / 100"
    )

    add_callout_box(
        doc,
        "TF-IDF cosine similarity measures text and vocabulary overlap between documents. "
        "It is NOT a probability of candidate job success, competency, or future job performance. "
        "'Skill Not Detected' signifies that the specific keyword/alias was not located in the digital text; "
        "it does NOT imply that the candidate definitely lacks that capability.",
        title="MANDATORY ETHICAL AI PRINCIPLES"
    )

    # ---------------------------------------------------------
    # 8. EXPERIMENTAL RESULTS & MULTI-ROLE SCREENING
    # ---------------------------------------------------------
    add_sec_heading("7. Experimental Results & Multi-Role Evaluation")
    doc.add_paragraph(
        "TalentLens AI was tested using seven synthetic demo documents representing diverse candidate profiles and technical edge cases, "
        "as well as real-world resume screening across distinct job descriptions."
    )

    res_tbl = doc.add_table(rows=8, cols=5)
    res_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_headers = ["Rank", "Candidate Document", "Overall Score", "Skill Coverage", "Extraction Status / Notes"]
    for idx, hname in enumerate(t_headers):
        c = res_tbl.rows[0].cells[idx]
        set_cell_background(c, "1E3A8A")
        set_cell_margins(c, 60, 60, 80, 80)
        p = c.paragraphs[0]
        r = p.add_run(hname)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

    demo_results = [
        ("#1", "candidate_01_aisha_sharma_senior_ml.pdf", "52.1 / 100", "94.1% (16/17 skills)", "Success — Strong ML/PyTorch match"),
        ("#2", "candidate_02_david_chen_data_analyst.pdf", "17.9 / 100", "35.3% (6/17 skills)", "Success — Moderate match; missing PyTorch/Docker"),
        ("#3", "candidate_04_marcus_vance_cloud_devops.pdf", "17.7 / 100", "35.3% (6/17 skills)", "Success — DevOps profile; missing deep learning"),
        ("#4", "candidate_03_elena_rodriguez_fullstack_dev.pdf", "14.5 / 100", "29.4% (5/17 skills)", "Success — Web developer; low ML overlap"),
        ("#5", "candidate_05_scanned_image_sample.pdf", "0.0 / 100", "0.0% (0 skills)", "Flagged: Scanned / Image-Only PDF (OCR Required)"),
        ("#6", "candidate_06_empty_sample.pdf", "0.0 / 100", "0.0% (0 skills)", "Flagged: Empty / Blank PDF (0 characters)"),
        ("#7", "candidate_07_corrupted_sample.pdf", "0.0 / 100", "0.0% (0 skills)", "Flagged: Corrupted or Invalid PDF stream"),
    ]
    for row_idx, data_tuple in enumerate(demo_results, start=1):
        row = res_tbl.rows[row_idx]
        bg = "F9FAFB" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data_tuple):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 50, 50, 70, 70)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(9)
            if col_idx == 0 or col_idx == 2:
                r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    chart_img_path = IMAGES_DIR / "role_comparison_chart.png"
    if chart_img_path.exists():
        doc.add_picture(str(chart_img_path), width=Inches(6.0))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figure 2: Dynamic Multi-Role Screening Verification on the Same Candidate Resume")
        r_cap.italic = True
        r_cap.font.size = Pt(9.5)
        r_cap.font.color.rgb = RGBColor(107, 114, 128)

    doc.add_paragraph(
        "Figure 2 illustrates the dynamic re-screening capability. When the same candidate resume is evaluated against an aligned role "
        "(AI Applications & Desktop Systems), the score increases to 49.1 / 100 with 100% skill match (9/9 skills matched). "
        "When evaluated against a Senior Machine Learning role requiring deep statistical modeling frameworks, the score drops to 9.6 / 100 "
        "reflecting genuine role misalignment."
    )

    # ---------------------------------------------------------
    # 9. QUALITY ASSURANCE & TESTING
    # ---------------------------------------------------------
    add_sec_heading("8. Quality Assurance & Automated Testing")
    doc.add_paragraph(
        "The project employs automated testing via pytest. A total of 30 comprehensive unit, integration, and UI workflow tests "
        "were developed and executed with a 100% pass rate."
    )

    test_img_path = IMAGES_DIR / "test_results_summary.png"
    if test_img_path.exists():
        doc.add_picture(str(test_img_path), width=Inches(6.0))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figure 3: Automated Test Execution Breakdown (30/30 Tests Passing)")
        r_cap.italic = True
        r_cap.font.size = Pt(9.5)
        r_cap.font.color.rgb = RGBColor(107, 114, 128)

    test_tbl = doc.add_table(rows=6, cols=3)
    test_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    th_headers = ["Test Suite / File", "Test Count", "Core Coverage & Scope"]
    for idx, h in enumerate(th_headers):
        c = test_tbl.rows[0].cells[idx]
        set_cell_background(c, "047857")
        set_cell_margins(c, 60, 60, 80, 80)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

    test_suites = [
        ("tests/test_pdf_extractor.py", "8 Tests", "PyMuPDF digital text extraction, whitespace normalization, corrupt/empty/scanned detection"),
        ("tests/test_skills_extractor.py", "8 Tests", "Multi-word phrase matching, symbol lookarounds (C++, C#, .NET), contextual aliases (pd, np)"),
        ("tests/test_matcher.py", "4 Tests", "TF-IDF cosine similarity, formula bounds clamping, top keyword contributions, candidate ranking"),
        ("tests/test_edge_cases.py", "5 Tests", "Empty job description handling, special characters, zero-match resumes, dynamic weight adjustments"),
        ("tests/test_multi_role_screening.py & test_ui_workflow.py", "5 Tests", "Multi-role rescreening, end-to-end batch screening, Streamlit AppTest UI simulation"),
    ]
    for row_idx, (tfile, tcount, tscope) in enumerate(test_suites, start=1):
        row = test_tbl.rows[row_idx]
        bg = "F9FAFB" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate([tfile, tcount, tscope]):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 50, 50, 70, 70)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(9)
            if col_idx == 1:
                r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ---------------------------------------------------------
    # 10. LIMITATIONS & FUTURE ENHANCEMENTS
    # ---------------------------------------------------------
    add_sec_heading("9. Limitations & Future Enhancements")
    add_sec_heading("9.1 System Limitations", level=2)
    doc.add_paragraph(
        "1. Absence of Optical Character Recognition (OCR): By design, TalentLens AI uses PyMuPDF for digital text streams. "
        "Scanned image resumes lacking digital text cannot be transcribed and are safely flagged.\n"
        "2. Raw TF-IDF Document Dilution: Cosine similarity between a compact job description and a lengthy multi-section resume "
        "naturally falls in the 15%–35% range due to differing document length and vocabulary breadth.\n"
        "3. Lexical Keyword Bounds: Skills extraction relies on predefined canonical terms and aliases; unlisted domain terminology "
        "must be added via the custom skill form."
    )

    add_sec_heading("9.2 Future Enhancements", level=2)
    doc.add_paragraph(
        "• Local Offline OCR: Integrate Tesseract or EasyOCR as an optional offline fallback for scanned physical documents.\n"
        "• Local Sentence-Transformer Embeddings: Supplement TF-IDF with lightweight local semantic embeddings (e.g., all-MiniLM-L6-v2) "
        "for conceptual paraphrase matching without external APIs.\n"
        "• Experience Duration Parser: Extract date spans to compute years of experience per detected technology.\n"
        "• Multi-Language Resume Ingestion: Expand taxonomy dictionaries to support non-English resumes."
    )

    # ---------------------------------------------------------
    # 11. CONCLUSION & REFERENCES
    # ---------------------------------------------------------
    add_sec_heading("10. Conclusion")
    doc.add_paragraph(
        "TalentLens AI demonstrates that automated candidate screening can be implemented locally, cost-effectively, and with complete "
        "algorithmic explainability. By replacing opaque black-box scoring with transparent mathematical formulas and combining "
        "PyMuPDF document extraction with robust TF-IDF and regex skill matching, the system resolves critical industry concerns regarding "
        "data privacy, algorithmic bias, and recruiter trust. All project requirements, unit tests, and edge-case guardrails have been "
        "verified through rigorous automated testing, delivering a production-grade academic capstone project."
    )

    add_sec_heading("11. References")
    refs = [
        "Salton, G., & Buckley, C. (1988). Term-weighting approaches in automatic text retrieval. Information Processing & Management, 24(5), 513-523.",
        "Manning, C. D., Raghavan, P., & Schütze, H. (2008). Introduction to Information Retrieval. Cambridge University Press.",
        "Pedregosa, F., et al. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825-2830.",
        "PyMuPDF Documentation. (2026). High performance PDF rendering and digital text extraction library. Artifex Software.",
        "Streamlit Inc. (2026). Streamlit: The fastest way to build and share data apps. Snowflake Inc.",
        "Barocas, S., Hardt, M., & Narayanan, A. (2023). Fairness and Machine Learning: Limitations and Opportunities. MIT Press."
    ]
    for idx, ref in enumerate(refs, start=1):
        p_ref = doc.add_paragraph()
        p_ref.paragraph_format.left_indent = Inches(0.4)
        p_ref.paragraph_format.first_line_indent = Inches(-0.4)
        r_num = p_ref.add_run(f"[{idx}] ")
        r_num.bold = True
        r_ref = p_ref.add_run(ref)
        r_ref.font.size = Pt(9.5)

    doc.save(str(DOCX_OUT))
    print(f"Successfully generated academic report at: {DOCX_OUT}")


if __name__ == "__main__":
    build_docx_report()
